from __future__ import annotations
# ==============================================================================
# ARCHIVO: database/repo.py
# PROPÓSITO: Repositorio central de datos y lógica de negocio para la Versión
#            Web en Streamlit. Gestiona la base de datos Excel, simulación
#            masiva de 50 estudiantes para los integrantes del grupo
#            (Alejandro Campo, Juan Esteban Cardona y Diego Fernando Gómez),
#            15 asignaturas por programa con múltiples grupos, horarios de
#            2, 3 o 4 horas sin clases en la franja de almuerzo (13:00 - 14:00),
#            validador de cruces de horario, control de cupos y generación de
#            PDFs y Excel en memoria (BytesIO) para descarga directa.
# ==============================================================================

from datetime import date
from io import BytesIO
from pathlib import Path
import random
import re
import unicodedata
import uuid

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image

BASE_DIR = Path(__file__).resolve().parent.parent
RUTA_EXCEL_WEB = BASE_DIR / "Registro_Universidad_Web.xlsx"
RUTA_LOGO_WEB = BASE_DIR / "logo_ucc.png"
RUTA_INSTALLER_EXE = BASE_DIR / "Instalador_Sistema_Universitario_UCC_Setup.exe"

CREDITOS_MINIMOS = 6
CREDITOS_MAXIMOS = 19
NOTA_APROBATORIA = 3.0

HORA_APERTURA = 6
HORA_CIERRE = 22
HORA_ALMUERZO_INICIO = 13
HORA_ALMUERZO_FIN = 14
DURACIONES_PERMITIDAS = (2, 3, 4)

ENCABEZADOS = {
    "Facultades": ["ID_Facultad", "Nombre_Facultad", "Decano"],
    "Programas": ["ID_Programa", "Nombre_Programa", "ID_Facultad", "Semestres"],
    "Materias": ["ID_Materia", "Nombre_Curso", "Creditos", "Cupos", "Docente", "ID_Programa", "Horario"],
    "Estudiantes": ["ID_Estudiante", "Nombre", "Apellidos", "Correo", "Telefono", "Semestre", "ID_Programa"],
    "Matriculas": ["ID_Matricula", "Cedula", "ID_Materia", "Creditos", "Periodo", "Fecha_Matricula", "Estado"],
    "Notas": ["ID_Nota", "ID_Matricula", "ID_Materia", "Nota1", "Nota2", "Nota3", "Definitiva", "Estado"],
}

MAPA_DIAS = {
    "lunes": "Lunes", "lun": "Lunes",
    "martes": "Martes", "mar": "Martes",
    "miercoles": "Miércoles", "mie": "Miércoles",
    "jueves": "Jueves", "jue": "Jueves",
    "viernes": "Viernes", "vie": "Viernes",
    "sabado": "Sábado", "sab": "Sábado",
}

FACULTADES_OFICIALES = [
    ["F01", "Facultad de Ingeniería", "Dr. Hernán Darío Gómez"],
    ["F02", "Facultad de Ciencias Económicas y Administrativas", "Dra. Claudia Patricia Mejía"],
    ["F03", "Facultad de Psicología y Ciencias Sociales", "Mg. Camilo Andrés Perea"],
    ["F04", "Facultad de Derecho y Ciencias Políticas", "Dra. María Fernanda Morales"],
    ["F05", "Facultad de Ciencias de la Salud", "Dr. Carlos Alberto Osorio"],
]

PROGRAMAS_OFICIALES = [
    ["P01", "Ingeniería de Software y Sistemas", "F01", 10],
    ["P02", "Ingeniería Industrial", "F01", 10],
    ["P03", "Administración de Empresas", "F02", 9],
    ["P04", "Contaduría Pública", "F02", 9],
    ["P05", "Psicología", "F03", 10],
    ["P06", "Derecho", "F04", 10],
    ["P07", "Enfermería", "F05", 8],
    ["P08", "Medicina", "F05", 12],
]

PREFIJOS_PROGRAMA = {
    "P01": "SIS", "P02": "IND", "P03": "ADM", "P04": "CON",
    "P05": "PSI", "P06": "DER", "P07": "ENF", "P08": "MED",
}

PLAN_ESTUDIOS = {
    "P01": [
        ("Fundamentos de Programación", 4), ("Cálculo Diferencial", 4),
        ("Lógica y Matemáticas Discretas", 3), ("Introducción a la Ingeniería de Software", 2),
        ("Programación Orientada a Objetos", 4), ("Cálculo Integral", 3),
        ("Estructuras de Datos y Algoritmos", 4), ("Bases de Datos Relacionales", 3),
        ("Física Mecánica y Laboratorio", 4), ("Ingeniería de Requisitos y UML", 3),
        ("Arquitectura y Patrones de Software", 3), ("Redes de Computadores", 3),
        ("Metodologías Ágiles y Scrum", 2), ("Inteligencia Artificial y Machine Learning", 4),
        ("Calidad y Pruebas de Software", 3),
    ],
    "P02": [
        ("Fundamentos de Ingeniería Industrial", 2), ("Cálculo Diferencial", 4),
        ("Química General", 3), ("Dibujo Técnico y CAD", 3),
        ("Física Mecánica", 4), ("Estadística Descriptiva", 3),
        ("Álgebra Lineal", 3), ("Procesos de Manufactura", 4),
        ("Probabilidad e Inferencia", 3), ("Investigación de Operaciones", 4),
        ("Ergonomía y Seguridad Industrial", 2), ("Gestión de la Calidad", 3),
        ("Logística y Cadena de Suministro", 3), ("Simulación de Procesos", 3),
        ("Gestión de Proyectos Industriales", 3),
    ],
    "P03": [
        ("Fundamentos de Administración", 3), ("Matemáticas Básicas", 3),
        ("Introducción a la Contabilidad", 3), ("Microeconomía", 3),
        ("Comunicación Empresarial", 2), ("Derecho Comercial", 2),
        ("Macroeconomía", 3), ("Estadística Aplicada", 3),
        ("Comportamiento Organizacional", 3), ("Matemáticas Financieras", 4),
        ("Gestión del Talento Humano", 3), ("Mercadeo Estratégico", 3),
        ("Finanzas Corporativas", 4), ("Emprendimiento e Innovación", 2),
        ("Planeación Estratégica", 3),
    ],
    "P04": [
        ("Contabilidad General", 4), ("Matemáticas Financieras", 3),
        ("Fundamentos de Economía", 3), ("Derecho Comercial y Laboral", 2),
        ("Costos y Presupuestos", 4), ("Estadística Descriptiva", 3),
        ("Contabilidad Intermedia", 4), ("Legislación Tributaria", 3),
        ("Sistemas de Información Contable", 3), ("Análisis Financiero", 3),
        ("Auditoría Financiera", 4), ("Revisoría Fiscal", 3),
        ("Contabilidad Pública", 3), ("NIIF y Normas Internacionales", 3),
        ("Ética Profesional del Contador", 2),
    ],
    "P05": [
        ("Introducción a la Psicología", 3), ("Procesos Psicológicos Básicos", 3),
        ("Neuroanatomía y Neurofisiología", 4), ("Epistemología de la Psicología", 2),
        ("Psicología del Desarrollo", 3), ("Estadística para Psicología", 3),
        ("Teorías de la Personalidad", 3), ("Psicología Social", 3),
        ("Psicopatología", 4), ("Evaluación Psicológica", 4),
        ("Metodología de la Investigación", 3), ("Psicología Educativa", 3),
        ("Psicología Clínica", 4), ("Psicología Organizacional", 3),
        ("Ética y Deontología Psicológica", 2),
    ],
    "P06": [
        ("Introducción al Derecho", 3), ("Teoría del Estado", 3),
        ("Historia del Derecho", 2), ("Derecho Romano", 3),
        ("Derecho Constitucional Colombiano", 4), ("Derecho Civil de Personas", 3),
        ("Derecho Penal General", 4), ("Teoría del Delito", 3),
        ("Derecho de Obligaciones", 3), ("Derecho Laboral", 3),
        ("Derecho Comercial", 3), ("Derecho Administrativo", 4),
        ("Derecho Procesal Civil", 4), ("Derechos Humanos y DIH", 2),
        ("Derecho Internacional Público", 3),
    ],
    "P07": [
        ("Fundamentos de Enfermería", 4), ("Anatomía Humana", 4),
        ("Biología Celular", 3), ("Química Orgánica", 3),
        ("Fisiología Humana", 4), ("Microbiología y Parasitología", 3),
        ("Bioquímica", 3), ("Cuidado Básico del Paciente", 4),
        ("Farmacología", 3), ("Salud Pública y Epidemiología", 3),
        ("Enfermería del Adulto", 4), ("Enfermería Materno Infantil", 4),
        ("Enfermería en Salud Mental", 3), ("Gerencia en Servicios de Salud", 2),
        ("Bioética en Enfermería", 2),
    ],
    "P08": [
        ("Biología Médica", 4), ("Anatomía General", 4),
        ("Histología y Embriología", 4), ("Química Médica", 3),
        ("Bioquímica Médica", 4), ("Fisiología Médica", 4),
        ("Microbiología Médica", 3), ("Inmunología", 3),
        ("Patología General", 4), ("Farmacología Médica", 4),
        ("Semiología", 4), ("Salud Pública", 3),
        ("Medicina Interna", 4), ("Pediatría", 4),
        ("Bioética Médica", 2),
    ],
}

_DOCENTES_NOMBRES = [
    ("Carlos Andrés", "M"), ("Ana María", "F"), ("Luis Fernando", "M"), ("Sandra Milena", "F"),
    ("Jorge Eliécer", "M"), ("Patricia", "F"), ("Mauricio", "M"), ("Marcela", "F"),
    ("Ricardo", "M"), ("Claudia", "F"), ("Andrés Felipe", "M"), ("Diana Carolina", "F"),
    ("Hernán", "M"), ("Martha Lucía", "F"), ("Roberto", "M"), ("Gloria Inés", "F"),
    ("Fernando", "M"), ("Liliana", "F"), ("Camilo", "M"), ("Adriana", "F"),
    ("Julián", "M"), ("Paola Andrea", "F"), ("Sergio", "M"), ("Natalia", "F"),
    ("Alexander", "M"), ("Carolina", "F"), ("Óscar", "M"), ("Mónica", "F"),
    ("Gustavo", "M"), ("Yolanda", "F"),
]
_DOCENTES_APELLIDOS = [
    "Pérez Mora", "Gómez Ríos", "Martínez Lenis", "Ortiz Valencia", "Henao Cardona",
    "Trujillo Sanz", "Pineda Cruz", "Lozano Reyes", "Echeverri Soto", "Beltrán Cano",
    "Valencia Mejía", "Caicedo Rojas", "Mina Angulo", "Rosero Paz", "Osorio Duque",
    "Quintero Arias", "Cifuentes Lara", "Zúñiga Mena", "Arboleda Vélez", "Ramírez Tello",
    "Salcedo Perea", "Barona Gil", "Lasso Ibarra", "Tascón Rey", "Sinisterra Lucumí",
]
_TITULOS = ["Ing.", "Mg.", "Esp.", "Lic.", "Dr."]

NOMBRES_POOL = [
    "Laura Sofía", "Mateo Andrés", "Valentina", "Daniel Eduardo", "Mariana",
    "Andrés Felipe", "Camila", "Sebastián", "Isabella", "Nicolás",
    "Gabriela", "Sara Lucía", "Samuel David", "Manuela", "Felipe",
    "Natalia", "Julián", "Paula Andrea", "Esteban", "Carolina",
    "Simón", "Valeria", "Daniela", "David", "Juliana",
    "Emmanuel", "Melissa", "Cristian", "Salomé", "Lucas",
    "Luciana", "Martín", "Vanessa", "Tomás", "Antonia",
    "Jerónimo", "María José", "Juan José", "Adriana", "Carlos Alberto",
    "Tatiana", "Oscar Iván", "Diana Marcela", "Camilo", "Ángela María"
]

APELLIDOS_POOL = [
    "Restrepo Henao", "Pelaez Castro", "Guzmán Osorio", "Vargas Silva", "Herrera Medina",
    "Salazar Duque", "Jaramillo Maya", "Giraldo Cano", "Montoya Zuluaga", "Bermúdez Cárdenas",
    "Londoño Ortiz", "Zapata Carvajal", "Quintero Blandón", "Álvarez Ospina", "Valencia Gallego",
    "Muñoz Agudelo", "Arango Bedoya", "Betancur Chaverra", "Escobar Forero", "Tamayo Villegas",
    "Caicedo Murillo", "Rojas Hurtado", "Castaño Hoyos", "Mejía Barrientos", "Lozano Mosquera",
    "Paredes Rivas", "Cifuentes Balanta", "Rosero Benavides", "Narváez Erazo", "Delgado Guerrero",
    "Chaves Zambrano", "Solano Peñaloza", "Bustos Cuellar", "Gutiérrez Peña", "Pineda Rueda"
]


def _nombre_docente(indice: int) -> str:
    nombre, genero = _DOCENTES_NOMBRES[indice % len(_DOCENTES_NOMBRES)]
    apellido = _DOCENTES_APELLIDOS[(indice * 7 + 3) % len(_DOCENTES_APELLIDOS)]
    titulo = _TITULOS[(indice // 3) % len(_TITULOS)]
    if titulo == "Dr." and genero == "F":
        titulo = "Dra."
    return f"{titulo} {nombre} {apellido}"


def inicios_validos(duracion: int) -> list[int]:
    """Horas de inicio válidas sin tocar la hora de almuerzo (13:00 - 14:00)."""
    dur = int(duracion)
    if dur not in DURACIONES_PERMITIDAS:
        return []
    horas = []
    for h in range(HORA_APERTURA, HORA_CIERRE - dur + 1):
        fin = h + dur
        if h < HORA_ALMUERZO_FIN and fin > HORA_ALMUERZO_INICIO:
            continue
        horas.append(h)
    return horas


def construir_horario(dia: str, hora_inicio: int, duracion: int) -> str:
    fin = int(hora_inicio) + int(duracion)
    return f"{dia} {int(hora_inicio):02d}:00 - {fin:02d}:00"


def codigo_base_materia(id_materia: str) -> str:
    texto = str(id_materia or "").strip().upper()
    return re.sub(r"-G\d+$", "", texto)


def construir_oferta_materias(semilla: int = 11) -> list[list]:
    azar = random.Random(semilla)
    dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
    oferta = []

    for idx_prog, programa in enumerate(PROGRAMAS_OFICIALES):
        id_programa = programa[0]
        prefijo = PREFIJOS_PROGRAMA[id_programa]

        for idx_curso, (nombre, creditos) in enumerate(PLAN_ESTUDIOS[id_programa]):
            codigo_base = f"{prefijo}{101 + idx_curso}"
            docente = _nombre_docente(idx_prog * 15 + idx_curso)
            validos = inicios_validos(creditos)
            mananas = [h for h in validos if h + creditos <= HORA_ALMUERZO_INICIO]
            tardes = [h for h in validos if h >= HORA_ALMUERZO_FIN]

            for grupo in (1, 2):
                bolsa = mananas if grupo == 1 else tardes
                inicio = azar.choice(bolsa or validos)
                dia = azar.choice(dias)
                docente_grupo = docente if grupo == 1 else _nombre_docente(idx_prog * 15 + idx_curso + 61)
                cupos = azar.choice([30, 32, 35, 40])
                oferta.append([
                    f"{codigo_base}-G{grupo}",
                    nombre,
                    creditos,
                    cupos,
                    docente_grupo,
                    id_programa,
                    construir_horario(dia, inicio, creditos),
                ])
    return oferta


# ------------------------------------------------------------------------------
# 1. INICIALIZACIÓN Y GENERADOR MASIVO
# ------------------------------------------------------------------------------
def _construir_libro_vacio() -> Workbook:
    """Construye un libro de Excel en memoria con las hojas y encabezados en blanco."""
    libro = Workbook()
    libro.remove(libro.active)

    for nombre_hoja, cols in ENCABEZADOS.items():
        hoja = libro.create_sheet(title=nombre_hoja)
        hoja.freeze_panes = "A2"
        for idx, col_titulo in enumerate(cols, start=1):
            c = hoja.cell(row=1, column=idx, value=col_titulo)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1F6AA5", end_color="1F6AA5", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            hoja.column_dimensions[get_column_letter(idx)].width = max(len(col_titulo) + 6, 16)
    return libro


def reiniciar_base_datos_vacia_web():
    """Vacía por completo la base de datos web dejándola en blanco (solo encabezados)."""
    libro = _construir_libro_vacio()
    libro.save(RUTA_EXCEL_WEB)
    libro.close()


def _seleccionar_materias_para_estudiante(ofertas_programa: list[list], azar: random.Random, objetivo_creditos: int = 15) -> list[list]:
    por_curso: dict[str, list[list]] = {}
    for m in ofertas_programa:
        if int(m[3]) <= 0:
            continue
        base = codigo_base_materia(m[0])
        por_curso.setdefault(base, []).append(m)

    codigos = list(por_curso.keys())
    azar.shuffle(codigos)

    seleccionadas: list[list] = []
    usadas_horario: list[dict] = []
    creditos = 0

    for base in codigos:
        grupos = list(por_curso[base])
        azar.shuffle(grupos)
        for cand in grupos:
            cr = int(cand[2])
            if creditos + cr > CREDITOS_MAXIMOS:
                continue
            info_cand = {"id": cand[0], "nombre": cand[1], "horario": cand[6]}
            hay_cruce, _ = validar_cruce_horarios(info_cand, usadas_horario)
            if hay_cruce:
                continue
            seleccionadas.append(cand)
            usadas_horario.append(info_cand)
            creditos += cr
            break
        if creditos >= objetivo_creditos:
            break
    return seleccionadas


def _construir_libro_simulacion_masiva(total_estudiantes: int = 50) -> tuple[Workbook, dict]:
    """Construye en memoria el libro con los 50 estudiantes y las 240 ofertas (15 por programa x 2 grupos)."""
    libro = _construir_libro_vacio()

    for f in FACULTADES_OFICIALES:
        libro["Facultades"].append(f)
    for p in PROGRAMAS_OFICIALES:
        libro["Programas"].append(p)

    oferta_materias = construir_oferta_materias(semilla=11)
    materias_dict = {m[0]: list(m) for m in oferta_materias}
    por_programa: dict[str, list[list]] = {}
    for m in materias_dict.values():
        por_programa.setdefault(str(m[5]), []).append(m)

    estudiantes = [
        ["1001", "Alejandro", "Campo Solarte", "alejandro.campo@campusucc.edu.co", "3158901234", 6, "P01"],
        ["1002", "Juan Esteban", "Cardona Calambas", "juan.cardona@campusucc.edu.co", "3176543210", 6, "P01"],
        ["1003", "Diego Fernando", "Gómez Rivera", "diego.gomez@campusucc.edu.co", "3201122334", 6, "P01"],
    ]

    programas_cods = [p[0] for p in PROGRAMAS_OFICIALES]
    azar = random.Random(42)

    for i in range(4, total_estudiantes + 1):
        cedula = f"{1000 + i}"
        nom = NOMBRES_POOL[(i - 4) % len(NOMBRES_POOL)]
        ape = APELLIDOS_POOL[(i - 4) % len(APELLIDOS_POOL)]
        prog = programas_cods[(i - 1) % len(programas_cods)]
        max_sem = next((p[3] for p in PROGRAMAS_OFICIALES if p[0] == prog), 10)
        sem = ((i * 3) % max_sem) + 1
        nom_l = _quitar_acentos(nom.lower().split()[0])
        ape_l = _quitar_acentos(ape.lower().split()[0])
        correo = f"{nom_l}.{ape_l}{i}@campusucc.edu.co"
        tel = f"31{azar.randint(0, 9)}{azar.randint(1000000, 9999999)}"
        estudiantes.append([cedula, nom, ape, correo, tel, sem, prog])

    for e in estudiantes:
        libro["Estudiantes"].append(e)

    matriculas = []
    notas = []
    periodo = "2026-2"

    notas_int = [
        [4.8, 4.6, 5.0],
        [4.5, 4.7, 4.9],
        [4.6, 4.8, 4.7],
    ]

    for idx_i, est_i in enumerate(estudiantes[:3]):
        ced_i = est_i[0]
        prog_i = est_i[6]
        elegidas = _seleccionar_materias_para_estudiante(por_programa.get(prog_i, []), azar, objetivo_creditos=15)
        for idx_m, m_info in enumerate(elegidas):
            cod_m = m_info[0]
            id_mat = f"MAT{ced_i}_{cod_m}"
            m_info[3] = max(0, int(m_info[3]) - 1)
            matriculas.append([id_mat, ced_i, cod_m, int(m_info[2]), periodo, "2026-09-10", "Activa"])

            n_b = notas_int[idx_i]
            v = (idx_m * 0.1) - 0.1
            n1 = min(5.0, max(3.5, round(n_b[0] + v, 2)))
            n2 = min(5.0, max(3.5, round(n_b[1] - v, 2)))
            n3 = min(5.0, max(3.5, round(n_b[2] + (v / 2), 2)))
            def_n = round((n1 + n2 + n3) / 3.0, 2)
            notas.append([f"NOT_{id_mat}", id_mat, cod_m, n1, n2, n3, def_n, "Aprobado"])

    for e in estudiantes[3:]:
        ced = e[0]
        prog = e[6]
        objetivo = azar.choice([12, 14, 15, 16])
        elegidas = _seleccionar_materias_para_estudiante(por_programa.get(prog, []), azar, objetivo_creditos=objetivo)
        for m_info in elegidas:
            cod_m = m_info[0]
            m_info[3] = max(0, int(m_info[3]) - 1)
            id_mat = f"MAT{ced}_{cod_m}"
            matriculas.append([id_mat, ced, cod_m, int(m_info[2]), periodo, "2026-09-15", "Activa"])

            if azar.random() < 0.85:
                n1 = round(azar.uniform(2.8, 4.9), 2)
                n2 = round(azar.uniform(2.7, 4.9), 2)
                n3 = round(azar.uniform(3.0, 5.0), 2)
                definitiva = round((n1 + n2 + n3) / 3.0, 2)
                estado = "Aprobado" if definitiva >= NOTA_APROBATORIA else "Reprobado"
                notas.append([f"NOT_{id_mat}", id_mat, cod_m, n1, n2, n3, definitiva, estado])

    for fila_m in materias_dict.values():
        libro["Materias"].append(fila_m)
    for mat in matriculas:
        libro["Matriculas"].append(mat)
    for n in notas:
        libro["Notas"].append(n)

    resumen = {
        "facultades": len(FACULTADES_OFICIALES),
        "programas": len(PROGRAMAS_OFICIALES),
        "materias": len(oferta_materias),
        "estudiantes": len(estudiantes),
        "matriculas": len(matriculas),
        "notas": len(notas),
    }
    return libro, resumen


def generar_excel_simulacion_bytes(total_estudiantes: int = 50) -> bytes:
    """
    Genera en memoria los bytes del archivo Excel con los 50 estudiantes y 240 grupos
    SIN sobrescribir la base de datos activa, para que el usuario pueda
    guardarlo en su PC e importarlo manualmente.
    """
    libro, _ = _construir_libro_simulacion_masiva(total_estudiantes)
    buf = BytesIO()
    libro.save(buf)
    libro.close()
    return buf.getvalue()


def importar_excel_subido_web(archivo_bytes: bytes) -> dict:
    """
    Importa un archivo Excel subido por el usuario (ej. Universidad_50_Estudiantes_UCC.xlsx)
    fusionando sus hojas en la base de datos activa sin duplicar llaves primarias.
    """
    asegurar_base_datos()
    libro_origen = load_workbook(BytesIO(archivo_bytes), data_only=True)
    libro_destino = load_workbook(RUTA_EXCEL_WEB)

    insertados = {h: 0 for h in ENCABEZADOS}
    duplicados = 0

    for nombre_hoja, cols in ENCABEZADOS.items():
        if nombre_hoja not in libro_origen.sheetnames:
            continue
        hoja_orig = libro_origen[nombre_hoja]
        hoja_dest = libro_destino[nombre_hoja]
        existentes = {
            str(f[0]).strip().lower()
            for f in hoja_dest.iter_rows(min_row=2, values_only=True)
            if f and f[0] is not None and str(f[0]).strip() != ""
        }
        for fila in hoja_orig.iter_rows(min_row=2, values_only=True):
            if not fila or fila[0] is None or str(fila[0]).strip() == "":
                continue
            llave = str(fila[0]).strip().lower()
            if llave in existentes:
                duplicados += 1
                continue
            hoja_dest.append(list(fila[:len(cols)]))
            existentes.add(llave)
            insertados[nombre_hoja] += 1

    libro_origen.close()
    libro_destino.save(RUTA_EXCEL_WEB)
    libro_destino.close()
    return {"insertados": insertados, "duplicados": duplicados}


def poblar_simulacion_masiva_web(total_estudiantes: int = 50) -> dict:
    libro, resumen = _construir_libro_simulacion_masiva(total_estudiantes)
    libro.save(RUTA_EXCEL_WEB)
    libro.close()
    return resumen


def asegurar_base_datos():
    """Crea la base de datos en blanco en la primera visita si no existe el archivo."""
    if not RUTA_EXCEL_WEB.exists():
        reiniciar_base_datos_vacia_web()


# ------------------------------------------------------------------------------
# 2. OPERACIONES CRUD SOBRE LA BASE DE DATOS
# ------------------------------------------------------------------------------
def obtener_datos(nombre_hoja: str) -> list[tuple]:
    """Devuelve todas las filas de `nombre_hoja` como lista de tuplas."""
    asegurar_base_datos()
    libro = load_workbook(RUTA_EXCEL_WEB, read_only=True, data_only=True)
    if nombre_hoja not in libro.sheetnames:
        libro.close()
        return []
    filas = [
        f for f in libro[nombre_hoja].iter_rows(min_row=2, values_only=True)
        if f and f[0] is not None and str(f[0]).strip() != ""
    ]
    libro.close()
    return filas


def buscar_fila(nombre_hoja: str, col_idx: int, valor):
    """Retorna la primera fila donde `fila[col_idx] == valor`."""
    buscado = str(valor).strip().lower()
    for fila in obtener_datos(nombre_hoja):
        if len(fila) > col_idx and fila[col_idx] is not None:
            if str(fila[col_idx]).strip().lower() == buscado:
                return fila
    return None


def insertar_fila(nombre_hoja: str, datos: list):
    """Agrega un nuevo registro al final de `nombre_hoja`."""
    asegurar_base_datos()
    libro = load_workbook(RUTA_EXCEL_WEB)
    libro[nombre_hoja].append(datos)
    libro.save(RUTA_EXCEL_WEB)
    libro.close()


def actualizar_fila(nombre_hoja: str, col_llave: int, valor_llave, nuevos_valores: list) -> bool:
    """Reemplaza los valores de la fila que coincida con `valor_llave`."""
    asegurar_base_datos()
    libro = load_workbook(RUTA_EXCEL_WEB)
    hoja = libro[nombre_hoja]
    buscado = str(valor_llave).strip().lower()
    actualizado = False

    for num_fila in range(2, hoja.max_row + 1):
        val_actual = hoja.cell(row=num_fila, column=col_llave + 1).value
        if val_actual is not None and str(val_actual).strip().lower() == buscado:
            for idx_col, val in enumerate(nuevos_valores, start=1):
                hoja.cell(row=num_fila, column=idx_col, value=val)
            actualizado = True
            break

    if actualizado:
        libro.save(RUTA_EXCEL_WEB)
    libro.close()
    return actualizado


# ------------------------------------------------------------------------------
# 3. ALGORITMO VALIDADOR DE CRUCES DE HORARIO Y GESTIÓN DE CUPOS
# ------------------------------------------------------------------------------
def _quitar_acentos(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def _hora_a_minutos(cadena_hora: str) -> int:
    limpia = cadena_hora.strip().lower().replace(" ", "")
    m = re.match(r"^(\d{1,2})(?::(\d{2}))?(am|pm|m)?$", limpia)
    if not m:
        raise ValueError(cadena_hora)
    h = int(m.group(1))
    mins = int(m.group(2)) if m.group(2) else 0
    suf = m.group(3)
    if suf == "pm" and h < 12:
        h += 12
    elif suf == "am" and h == 12:
        h = 0
    return h * 60 + mins


def parsear_horario(texto_horario: str) -> list[dict]:
    if not texto_horario:
        return []
    bloques = []
    for seg in re.split(r"[;/|]", str(texto_horario)):
        seg_norm = _quitar_acentos(seg.lower())
        dias = []
        for p in re.findall(r"[a-z]+", seg_norm):
            if p in MAPA_DIAS and MAPA_DIAS[p] not in dias:
                dias.append(MAPA_DIAS[p])
        rango = re.search(r"(\d{1,2}(?::\d{2})?\s*(?:am|pm|m)?)\s*[-a–]\s*(\d{1,2}(?::\d{2})?\s*(?:am|pm|m)?)", seg_norm)
        if dias and rango:
            try:
                ini = _hora_a_minutos(rango.group(1))
                fin = _hora_a_minutos(rango.group(2))
                if ini < 360 and "am" not in rango.group(1):
                    ini += 12 * 60
                if fin <= ini:
                    fin += 12 * 60
                for d in dias:
                    bloques.append({"dia": d, "inicio": ini, "fin": fin})
            except ValueError:
                continue
    return bloques


def validar_cruce_horarios(nueva_mat: dict, lista_existentes: list[dict]) -> tuple[bool, str]:
    bloques_nuevos = parsear_horario(nueva_mat.get("horario", ""))
    for ex in lista_existentes:
        for bn in bloques_nuevos:
            for be in parsear_horario(ex.get("horario", "")):
                if bn["dia"] == be["dia"]:
                    ini = max(bn["inicio"], be["inicio"])
                    fin = min(bn["fin"], be["fin"])
                    if ini < fin:
                        h_ini = f"{ini // 60:02d}:{ini % 60:02d}"
                        h_fin = f"{fin // 60:02d}:{fin % 60:02d}"
                        return True, (
                            f"Cruce detectado el {bn['dia']} ({h_ini} a {h_fin}): "
                            f"'{nueva_mat['nombre']}' choca con '{ex['nombre']}'."
                        )
    return False, ""


def ajustar_cupo_materia(id_materia: str, delta: int) -> bool:
    """Modifica el cupo de una asignatura sumando `delta` (-1 al matricular, +1 al cancelar)."""
    fila = buscar_fila("Materias", 0, id_materia)
    if not fila:
        return False
    cupo_actual = int(fila[3]) if fila[3] is not None else 0
    nuevo_cupo = max(0, cupo_actual + delta)
    return actualizar_fila(
        "Materias", 0, id_materia,
        [fila[0], fila[1], int(fila[2]), nuevo_cupo, fila[4], fila[5], fila[6]]
    )


def guardar_o_actualizar_nota(id_matricula: str, id_materia: str, n1: float, n2: float, n3: float) -> tuple[float, str]:
    """Inserta o actualiza las calificaciones de una matrícula sin duplicar filas."""
    definitiva = round((n1 + n2 + n3) / 3.0, 2)
    estado = "Aprobado" if definitiva >= NOTA_APROBATORIA else "Reprobado"
    existente = buscar_fila("Notas", 1, id_matricula)

    if existente:
        id_nota = str(existente[0])
        actualizar_fila("Notas", 1, id_matricula, [id_nota, id_matricula, id_materia, n1, n2, n3, definitiva, estado])
    else:
        id_nota = str(uuid.uuid4())[:8].upper()
        insertar_fila("Notas", [id_nota, id_matricula, id_materia, n1, n2, n3, definitiva, estado])

    return definitiva, estado


# ------------------------------------------------------------------------------
# 4. GENERADORES DE DOCUMENTOS PDF EN MEMORIA (BytesIO para st.download_button)
# ------------------------------------------------------------------------------
def generar_pdf_matricula_bytes(cedula: str) -> bytes | None:
    """Genera el PDF del Recibo de Matrícula en memoria listo para descargar."""
    fila_est = buscar_fila("Estudiantes", 0, cedula)
    if not fila_est:
        return None

    matriculas = [
        m for m in obtener_datos("Matriculas")
        if str(m[1]).strip() == str(cedula).strip() and (len(m) < 7 or str(m[6]).lower() == "activa")
    ]
    if not matriculas:
        return None

    prog = buscar_fila("Programas", 0, fila_est[6])
    nom_prog = prog[1] if prog else str(fila_est[6])
    mapa_mat = {str(m[0]): m for m in obtener_datos("Materias")}

    datos_tabla = [["Código", "Asignatura", "Docente", "Horario", "Créd.", "Periodo"]]
    total_cr = 0
    for m in matriculas:
        info = mapa_mat.get(str(m[2]))
        nom_c = info[1] if info else str(m[2])
        doc_c = info[4] if info else "Por asignar"
        hor_c = info[6] if info else "Por definir"
        cr = int(m[3])
        total_cr += cr
        datos_tabla.append([str(m[2]), nom_c, doc_c, hor_c, str(cr), str(m[4])])

    datos_tabla.append(["", "TOTAL CRÉDITOS MATRICULADOS", "", "", str(total_cr), ""])

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=1.8*cm, leftMargin=1.8*cm, topMargin=1.8*cm, bottomMargin=1.8*cm)
    estilos = getSampleStyleSheet()
    estilo_tit = ParagraphStyle("TitWeb", parent=estilos["Title"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#1F6AA5"), alignment=TA_CENTER)
    estilo_pie = ParagraphStyle("PieWeb", parent=estilos["Normal"], fontName="Helvetica-Oblique", fontSize=8, textColor=colors.grey, alignment=TA_RIGHT)

    elems = []
    if RUTA_LOGO_WEB.exists():
        img = Image(str(RUTA_LOGO_WEB), width=2.6*cm, height=2.6*cm)
        img.hAlign = "CENTER"
        elems.append(img)
        elems.append(Spacer(1, 0.25*cm))

    elems.append(Paragraph("UNIVERSIDAD COOPERATIVA DE COLOMBIA — SEDE CALI", estilo_tit))
    elems.append(Spacer(1, 0.3*cm))
    elems.append(Paragraph(
        f"<b>Estudiante:</b> {fila_est[1]} {fila_est[2]} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Documento:</b> {cedula}<br/>"
        f"<b>Programa:</b> {nom_prog} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Semestre:</b> {fila_est[5]} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Fecha:</b> {date.today().strftime('%Y-%m-%d')}",
        estilos["Normal"]
    ))
    elems.append(Spacer(1, 0.5*cm))

    tabla = Table(datos_tabla, colWidths=[2.0*cm, 4.8*cm, 3.6*cm, 4.4*cm, 1.5*cm, 1.8*cm])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E8B57")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E8F5E9")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    elems.append(tabla)
    elems.append(Spacer(1, 0.6*cm))
    elems.append(Paragraph(f"Comprobante Oficial Verificado — Código: {str(uuid.uuid4()).upper()[:10]}", estilo_pie))

    doc.build(elems)
    return buffer.getvalue()


def generar_pdf_boletin_bytes(cedula: str) -> bytes | None:
    """Genera el Boletín Oficial de Calificaciones en memoria listo para descargar."""
    fila_est = buscar_fila("Estudiantes", 0, cedula)
    if not fila_est:
        return None

    matriculas = [
        m for m in obtener_datos("Matriculas")
        if str(m[1]).strip() == str(cedula).strip() and (len(m) < 7 or str(m[6]).lower() == "activa")
    ]
    if not matriculas:
        return None

    prog = buscar_fila("Programas", 0, fila_est[6])
    nom_prog = prog[1] if prog else str(fila_est[6])
    mapa_mat = {str(m[0]): m for m in obtener_datos("Materias")}
    mapa_notas = {str(n[1]): n for n in obtener_datos("Notas")}

    datos_tabla = [["Código", "Asignatura", "Créd.", "Nota 1", "Nota 2", "Nota 3", "Definitiva", "Estado"]]
    suma_pond = 0.0
    cred_cal = 0

    for m in matriculas:
        id_mat = str(m[0])
        cod_c = str(m[2])
        cr = int(m[3])
        info_c = mapa_mat.get(cod_c)
        nom_c = info_c[1] if info_c else cod_c
        nota = mapa_notas.get(id_mat)

        if nota:
            n1, n2, n3, def_v, est_n = f"{float(nota[3]):.2f}", f"{float(nota[4]):.2f}", f"{float(nota[5]):.2f}", float(nota[6]), str(nota[7])
            suma_pond += def_v * cr
            cred_cal += cr
            def_s = f"{def_v:.2f}"
        else:
            n1, n2, n3, def_s, est_n = "-", "-", "-", "-", "Sin calificar"

        datos_tabla.append([cod_c, nom_c, str(cr), n1, n2, n3, def_s, est_n])

    pps = round(suma_pond / cred_cal, 2) if cred_cal > 0 else 0.0
    datos_tabla.append(["", "PROMEDIO PONDERADO SEMESTRAL (PPS)", str(cred_cal), "", "", "", f"{pps:.2f}", ""])

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=1.8*cm, leftMargin=1.8*cm, topMargin=1.8*cm, bottomMargin=1.8*cm)
    estilos = getSampleStyleSheet()
    estilo_tit = ParagraphStyle("TitBolWeb", parent=estilos["Title"], fontName="Helvetica-Bold", fontSize=15, textColor=colors.HexColor("#1F6AA5"), alignment=TA_CENTER)

    elems = []
    if RUTA_LOGO_WEB.exists():
        img = Image(str(RUTA_LOGO_WEB), width=2.6*cm, height=2.6*cm)
        img.hAlign = "CENTER"
        elems.append(img)
        elems.append(Spacer(1, 0.25*cm))

    elems.append(Paragraph("UNIVERSIDAD COOPERATIVA DE COLOMBIA — SEDE CALI", estilo_tit))
    elems.append(Spacer(1, 0.3*cm))
    elems.append(Paragraph(
        f"<b>Boletín Oficial de Calificaciones</b><br/>"
        f"<b>Estudiante:</b> {fila_est[1]} {fila_est[2]} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Documento:</b> {cedula}<br/>"
        f"<b>Programa:</b> {nom_prog} &nbsp;&nbsp;|&nbsp;&nbsp; <b>PPS:</b> {pps:.2f} / 5.00",
        estilos["Normal"]
    ))
    elems.append(Spacer(1, 0.5*cm))

    tabla = Table(datos_tabla, colWidths=[2.0*cm, 5.5*cm, 1.5*cm, 1.6*cm, 1.6*cm, 1.6*cm, 2.0*cm, 2.4*cm])
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F6AA5")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E3F2FD")),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ]))
    elems.append(tabla)
    doc.build(elems)
    return buffer.getvalue()


def exportar_excel_bytes() -> bytes:
    """Devuelve el archivo Excel completo en memoria para descarga directa."""
    asegurar_base_datos()
    return RUTA_EXCEL_WEB.read_bytes()


def generar_zip_escritorio_bytes() -> bytes | None:
    """Empaqueta dinámicamente la carpeta de la versión de escritorio en un archivo ZIP listo para usar."""
    import zipfile
    ruta_desk = BASE_DIR.parent / "1_Version_Escritorio_Avanzada"
    if not ruta_desk.exists():
        return None
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for f in ruta_desk.rglob("*"):
            if "__pycache__" in f.parts or f.suffix == ".pyc" or f.is_dir():
                continue
            rel = f.relative_to(ruta_desk)
            zf.write(f, arcname=str(Path("1_Version_Escritorio_Avanzada") / rel))
        bat_raiz = BASE_DIR.parent / "Ejecutar_Escritorio.bat"
        if bat_raiz.exists():
            zf.write(bat_raiz, arcname="Ejecutar_Escritorio.bat")
        req_raiz = BASE_DIR.parent / "requirements.txt"
        if req_raiz.exists():
            zf.write(req_raiz, arcname="requirements.txt")
    buf.seek(0)
    return buf.getvalue()


def obtener_instalador_exe_bytes() -> bytes | None:
    """Devuelve los bytes del instalador oficial .exe para Windows (Wizard con siguiente y acceso directo)."""
    if RUTA_INSTALLER_EXE.exists():
        return RUTA_INSTALLER_EXE.read_bytes()
    alt = BASE_DIR.parent / "dist_installer" / "Instalador_Sistema_Universitario_UCC_Setup.exe"
    if alt.exists():
        return alt.read_bytes()
    return None
