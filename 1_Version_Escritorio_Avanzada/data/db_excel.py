from __future__ import annotations
# ==============================================================================
# ARCHIVO: data/db_excel.py
# PROPÓSITO: Motor de persistencia relacional sobre Excel (.xlsx) usando OpenPyXL.
#            Incluye creación automática de esquema en blanco, reinicio/vaciado
#            total de base de datos, formato corporativo, filtros, inmovilización
#            de paneles y operaciones CRUD completas (Crear, Leer, Actualizar,
#            Eliminar y Filtrar registros sin duplicidades).
# ==============================================================================

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import config

# Alias para mantener compatibilidad limpia con todos los módulos
ruta_archivo = config.RUTA_ARCHIVO_EXCEL
encabezados = config.ENCABEZADOS_BD


def crear_archivo_excel():
    """
    Verifica si el archivo Excel de base de datos existe.
    Si no existe, crea todas las hojas definidas en `config.ENCABEZADOS_BD`
    en blanco (solo encabezados con diseño institucional UCC).
    Si ya existe, verifica que todas las hojas requeridas estén presentes.
    """
    borde_delgado = Side(border_style="thin", color="D0D7DE")
    borde_celda = Border(left=borde_delgado, right=borde_delgado, top=borde_delgado, bottom=borde_delgado)

    if ruta_archivo.exists():
        libro = load_workbook(ruta_archivo)
        modificado = False
        for nombre_hoja, columnas in encabezados.items():
            if nombre_hoja not in libro.sheetnames:
                hoja = libro.create_sheet(nombre_hoja)
                _formatear_encabezado_hoja(hoja, columnas, borde_celda)
                modificado = True
            else:
                # Actualiza títulos de columnas si el esquema cambió (Ej: Materias -> ID_Programa)
                hoja_existente = libro[nombre_hoja]
                for col_idx, titulo in enumerate(columnas, start=1):
                    if hoja_existente.cell(row=1, column=col_idx).value != titulo:
                        hoja_existente.cell(row=1, column=col_idx, value=titulo)
                        modificado = True
        if modificado:
            libro.save(ruta_archivo)
        return libro

    return reiniciar_base_de_datos_vacia()


def reiniciar_base_de_datos_vacia():
    """
    Elimina todos los registros y recrea el archivo `Registro_Matricula_Pro.xlsx`
    completamente en blanco, conservando únicamente las hojas y sus encabezados
    formateados con la identidad visual UCC.
    """
    borde_delgado = Side(border_style="thin", color="D0D7DE")
    borde_celda = Border(left=borde_delgado, right=borde_delgado, top=borde_delgado, bottom=borde_delgado)

    libro = Workbook()
    libro.remove(libro.active)  # Removemos la hoja por defecto

    for nombre_hoja, columnas in encabezados.items():
        hoja = libro.create_sheet(title=nombre_hoja)
        _formatear_encabezado_hoja(hoja, columnas, borde_celda)

    libro.save(ruta_archivo)
    return libro


def _formatear_encabezado_hoja(hoja, columnas: list[str], borde_celda):
    """Aplica el estilo azul corporativo UCC, filtros y congelación a la fila 1."""
    hoja.freeze_panes = "A2"
    ultima_letra = get_column_letter(len(columnas))
    hoja.auto_filter.ref = f"A1:{ultima_letra}1"
    hoja.row_dimensions[1].height = 24

    for col_idx, titulo in enumerate(columnas, start=1):
        celda = hoja.cell(row=1, column=col_idx, value=titulo)
        celda.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color="1F6AA5", end_color="1F6AA5", fill_type="solid")
        celda.alignment = Alignment(horizontal="center", vertical="center")
        celda.border = borde_celda
        letra_col = get_column_letter(col_idx)
        hoja.column_dimensions[letra_col].width = max(len(titulo) + 6, 16)


def cargar_libro(solo_lectura: bool = False):
    """Abre y retorna el objeto Workbook asegurando que exista previamente."""
    if not ruta_archivo.exists():
        crear_archivo_excel()
    return load_workbook(ruta_archivo, read_only=solo_lectura, data_only=True)


def guardar_datos(hoja, datos: list):
    """Agrega una nueva fila al final de la hoja especificada y guarda el libro."""
    hoja.append(datos)
    hoja.parent.save(ruta_archivo)


def insertar_en_hoja(nombre_hoja: str, datos: list):
    """Abre el libro, inserta una fila en `nombre_hoja` y guarda de forma segura."""
    libro = cargar_libro()
    hoja = libro[nombre_hoja]
    hoja.append(datos)
    libro.save(ruta_archivo)
    libro.close()


def buscar_fila(hoja, columna: int, valor):
    """
    Busca la primera fila cuyo valor en `columna` (índice base 0) coincida con `valor`.
    Retorna la tupla de valores de la fila o None si no existe.
    """
    buscado = str(valor).strip().lower()
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila and fila[columna] is not None:
            if str(fila[columna]).strip().lower() == buscado:
                return fila
    return None


def buscar_filas(hoja, columna: int, valor) -> list[tuple]:
    """
    Busca TODAS las filas cuyo valor en `columna` (índice base 0) coincida con `valor`.
    Retorna una lista de tuplas.
    """
    buscado = str(valor).strip().lower()
    resultados = []
    for fila in hoja.iter_rows(min_row=2, values_only=True):
        if fila and fila[columna] is not None:
            if str(fila[columna]).strip().lower() == buscado:
                resultados.append(fila)
    return resultados


def actualizar_fila(nombre_hoja: str, col_llave: int, valor_llave, nuevos_valores: list) -> bool:
    """
    Busca en `nombre_hoja` la fila donde la columna `col_llave` (base 0) sea igual a
    `valor_llave` y reemplaza todas sus celdas por `nuevos_valores`.
    Retorna True si se encontró y actualizó el registro, o False en caso contrario.
    """
    libro = cargar_libro()
    if nombre_hoja not in libro.sheetnames:
        libro.close()
        return False

    hoja = libro[nombre_hoja]
    buscado = str(valor_llave).strip().lower()
    encontrado = False

    for num_fila in range(2, hoja.max_row + 1):
        valor_celda = hoja.cell(row=num_fila, column=col_llave + 1).value
        if valor_celda is not None and str(valor_celda).strip().lower() == buscado:
            for idx_col, nuevo_val in enumerate(nuevos_valores, start=1):
                hoja.cell(row=num_fila, column=idx_col, value=nuevo_val)
            encontrado = True
            break

    if encontrado:
        libro.save(ruta_archivo)
    libro.close()
    return encontrado


def eliminar_fila(nombre_hoja: str, col_llave: int, valor_llave) -> bool:
    """
    Elimina de `nombre_hoja` la fila cuya columna `col_llave` (base 0) coincida
    con `valor_llave`. Retorna True si fue eliminada.
    """
    libro = cargar_libro()
    if nombre_hoja not in libro.sheetnames:
        libro.close()
        return False

    hoja = libro[nombre_hoja]
    buscado = str(valor_llave).strip().lower()
    eliminado = False

    for num_fila in range(2, hoja.max_row + 1):
        valor_celda = hoja.cell(row=num_fila, column=col_llave + 1).value
        if valor_celda is not None and str(valor_celda).strip().lower() == buscado:
            hoja.delete_rows(num_fila, 1)
            eliminado = True
            break

    if eliminado:
        libro.save(ruta_archivo)
    libro.close()
    return eliminado


def obtener_datos(nombre_hoja: str) -> list[tuple]:
    """Devuelve todas las filas no vacías de una hoja (omitiendo el encabezado)."""
    try:
        libro = cargar_libro(solo_lectura=True)
        if nombre_hoja not in libro.sheetnames:
            libro.close()
            return []
        filas = [
            fila for fila in libro[nombre_hoja].iter_rows(min_row=2, values_only=True)
            if fila and fila[0] is not None and str(fila[0]).strip() != ""
        ]
        libro.close()
        return filas
    except Exception as error:
        print(f"Error al leer la hoja '{nombre_hoja}': {error}")
        return []


def obtener_opciones(nombre_hoja: str, col_id: int, col_nombre: int) -> list[str]:
    """
    Construye una lista de cadenas con formato 'ID - Nombre' para poblar
    listas desplegables (CTkComboBox).
    """
    filas = obtener_datos(nombre_hoja)
    opciones = []
    for fila in filas:
        if len(fila) > max(col_id, col_nombre) and fila[col_id] is not None:
            opciones.append(f"{str(fila[col_id]).strip()} - {str(fila[col_nombre]).strip()}")
    return opciones


def extraer_id(seleccion_combobox: str) -> str:
    """Extrae el identificador limpio de una cadena 'ID - Nombre'."""
    if not seleccion_combobox:
        return ""
    return str(seleccion_combobox).split(" - ")[0].strip()


def obtener_estadisticas_sistema() -> dict:
    """Calcula indicadores clave (KPIs) en tiempo real para el panel de control."""
    facultades = obtener_datos("Facultades")
    programas = obtener_datos("Programas")
    materias = obtener_datos("Materias")
    estudiantes = obtener_datos("Estudiantes")
    matriculas = [
        m for m in obtener_datos("Matriculas")
        if len(m) < 7 or str(m[6]).strip().lower() == "activa"
    ]
    notas = obtener_datos("Notas")

    promedios = []
    aprobados = 0
    for n in notas:
        try:
            definitiva = float(n[6])
            promedios.append(definitiva)
            if definitiva >= config.NOTA_APROBATORIA:
                aprobados += 1
        except (ValueError, TypeError, IndexError):
            continue

    promedio_global = round(sum(promedios) / len(promedios), 2) if promedios else 0.0
    tasa_aprobacion = round((aprobados / len(promedios)) * 100, 1) if promedios else 0.0

    return {
        "total_facultades": len(facultades),
        "total_programas": len(programas),
        "total_materias": len(materias),
        "total_estudiantes": len(estudiantes),
        "matriculas_activas": len(matriculas),
        "total_calificaciones": len(notas),
        "promedio_global": promedio_global,
        "tasa_aprobacion": tasa_aprobacion,
    }



# ------------------------------------------------------------------------------
# CONSULTAS RELACIONALES AUXILIARES (Estudiante -> Programa -> Asignaturas)
# ------------------------------------------------------------------------------
def obtener_estudiante(cedula: str):
    """Retorna la fila completa del estudiante con documento `cedula` o None si no existe."""
    buscado = str(cedula).strip()
    for fila in obtener_datos("Estudiantes"):
        if str(fila[0]).strip() == buscado:
            return fila
    return None


def obtener_nombre_programa(id_programa: str) -> str:
    """Devuelve el nombre del programa académico a partir de su código."""
    buscado = str(id_programa).strip()
    for fila in obtener_datos("Programas"):
        if str(fila[0]).strip() == buscado:
            return str(fila[1]).strip()
    return buscado


def obtener_materias_programa(id_programa: str) -> list[tuple]:
    """
    Devuelve únicamente las asignaturas (grupos) que pertenecen al programa
    académico indicado. Cada programa tiene su propio plan de asignaturas.
    """
    buscado = str(id_programa).strip().lower()
    return [
        fila for fila in obtener_datos("Materias")
        if len(fila) > 5 and fila[5] is not None and str(fila[5]).strip().lower() == buscado
    ]


def obtener_matriculas_activas(cedula: str) -> list[tuple]:
    """Retorna las matrículas activas de un estudiante."""
    buscado = str(cedula).strip()
    return [
        fila for fila in obtener_datos("Matriculas")
        if str(fila[1]).strip() == buscado
        and (len(fila) < 7 or str(fila[6]).strip().lower() == "activa")
    ]
