from __future__ import annotations
# ==============================================================================
# ARCHIVO: poblar_datos.py
# PROPÓSITO: Generador de Archivos Excel de Simulación Masiva para la Universidad
#            Cooperativa de Colombia (UCC - Sede Cali).
#
#            Integrantes del proyecto:
#            • Alejandro Campo Solarte (Cédula 1001)
#            • Juan Esteban Cardona Calambas (Cédula 1002)
#            • Diego Fernando Gómez Rivera (Cédula 1003)
#
#            Genera un archivo Excel completo con:
#            • 5 Facultades y 8 Programas Académicos.
#            • 120 asignaturas (15 por programa) ofrecidas en 2 grupos con
#              horarios distintos (240 filas en la hoja 'Materias').
#            • 50 estudiantes que solo se matriculan en asignaturas de SU programa,
#              sin cruces de horario, respetando cupos y la hora de almuerzo.
#            • Matrículas, notas y cuentas de usuario.
#            El archivo se guarda donde el usuario elija (Explorador de Archivos)
#            para luego importarlo manualmente.
# ==============================================================================

from pathlib import Path
import random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import data.db_excel as db
from data.catalogo_academico import (
    FACULTADES_OFICIALES,
    PROGRAMAS_OFICIALES,
    PLAN_ESTUDIOS,
    construir_oferta_materias,
)
from core.seguridad import encriptar_password
from core.validador_horarios import verificar_cruce_entre_dos, codigo_base_materia

NOMBRES_POOL = [
    "Laura Sofía", "Mateo Andrés", "Valentina", "Daniel Eduardo", "Mariana",
    "Andrés Felipe", "Camila", "Sebastián", "Isabella", "Nicolás",
    "Gabriela", "Sara Lucía", "Samuel David", "Manuela", "Felipe",
    "Natalia", "Julián", "Paula Andrea", "Esteban", "Carolina",
    "Simón", "Valeria", "Daniela", "David", "Juliana",
    "Emmanuel", "Melissa", "Cristian", "Salomé", "Lucas",
    "Luciana", "Martín", "Vanessa", "Tomás", "Antonia",
    "Jerónimo", "María José", "Juan José", "Adriana", "Carlos Alberto",
    "Tatiana", "Oscar Iván", "Diana Marcela", "Camilo", "Ángela María",
    "Santiago", "Brayan Alexis", "Yuliana"
]

APELLIDOS_POOL = [
    "Restrepo Henao", "Pelaez Castro", "Guzmán Osorio", "Vargas Silva", "Herrera Medina",
    "Salazar Duque", "Jaramillo Maya", "Giraldo Cano", "Montoya Zuluaga", "Bermúdez Cárdenas",
    "Londoño Ortiz", "Zapata Carvajal", "Quintero Blandón", "Álvarez Ospina", "Valencia Gallego",
    "Muñoz Agudelo", "Arango Bedoya", "Betancur Chaverra", "Escobar Forero", "Tamayo Villegas",
    "Caicedo Murillo", "Rojas Hurtado", "Castaño Hoyos", "Mejía Barrientos", "Lozano Mosquera",
    "Paredes Rivas", "Cifuentes Balanta", "Rosero Benavides", "Narváez Erazo", "Delgado Guerrero",
    "Chaves Zambrano", "Solano Peñaloza", "Bustos Cuellar", "Gutiérrez Peña", "Pineda Rueda",
    "Sarmiento Beltrán", "Pinzón Rangel", "Pardo Téllez", "Orjuela Forero", "Mendoza Salamanca",
    "Barrera Nieto", "Cárdenas Prieto", "Suárez Garzón", "Navarro Pabón", "Villamizar Cáceres",
    "Carrillo Flórez", "Morales Ríos", "Mosquera Hinestroza"
]


def _quitar_tildes(texto: str) -> str:
    """Reemplaza vocales con tilde para armar correos institucionales válidos."""
    for origen, destino in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n")):
        texto = texto.replace(origen, destino)
    return texto


def _elegir_grupos_sin_cruce(
    programa: str,
    semestre: int,
    oferta_por_curso: dict,
    azar: random.Random,
) -> list[list]:
    """
    Selecciona para un estudiante los grupos de sus asignaturas del periodo
    (5 asignaturas del bloque que corresponde a su semestre) garantizando:
      - Solo asignaturas de su programa.
      - Ningún cruce de horario entre los grupos elegidos.
      - Cupos disponibles en el grupo escogido.
    Retorna la lista de filas de materias (grupos) elegidas.
    """
    plan = PLAN_ESTUDIOS[programa]
    bloque = (semestre - 1) % 3          # 3 bloques de 5 asignaturas por programa
    indices = range(bloque * 5, bloque * 5 + 5)

    elegidas: list[list] = []
    for idx in indices:
        nombre_curso = plan[idx][0]
        grupos = list(oferta_por_curso[(programa, nombre_curso)])
        azar.shuffle(grupos)
        for grupo in grupos:
            if grupo[3] <= 0:
                continue
            if any(verificar_cruce_entre_dos(grupo[6], e[6])[0] for e in elegidas):
                continue
            elegidas.append(grupo)
            break
    return elegidas


def generar_simulacion_masiva_universidad(
    total_estudiantes: int = 50,
    silencioso: bool = False,
    ruta_destino: Path | str | None = None
) -> dict:
    """
    Construye un archivo Excel completo con una universidad poblada:
    1. Registra a los 3 integrantes del grupo en las primeras 3 posiciones (1001, 1002, 1003).
    2. Registra estudiantes adicionales distribuidos en los 8 programas académicos.
    3. Crea 15 asignaturas por programa, cada una en 2 grupos con horarios distintos.
    4. Matricula a cada estudiante solo en materias de su programa, sin cruces de horario.
    5. Descuenta los cupos reales de cada grupo en la hoja 'Materias'.
    6. Registra calificaciones con definitivas calculadas y cuentas de usuario.
    7. Guarda el archivo en `ruta_destino` (si el usuario eligió dónde exportarlo para luego
       importarlo manualmente) o en `db.ruta_archivo` si no se especifica ruta.
    """
    archivo_salida = Path(ruta_destino) if ruta_destino else db.ruta_archivo
    azar = random.Random(42)

    libro = Workbook()
    libro.remove(libro.active)

    # 1. Crear hojas con encabezados institucionales
    for nombre_hoja, cols in config.ENCABEZADOS_BD.items():
        hoja = libro.create_sheet(title=nombre_hoja)
        hoja.freeze_panes = "A2"
        hoja.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
        for idx, titulo in enumerate(cols, start=1):
            c = hoja.cell(row=1, column=idx, value=titulo)
            c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="1F6AA5", end_color="1F6AA5", fill_type="solid")
            c.alignment = Alignment(horizontal="center", vertical="center")
            hoja.column_dimensions[get_column_letter(idx)].width = max(len(titulo) + 6, 16)

    # 2. Facultades y Programas
    for f in FACULTADES_OFICIALES:
        libro["Facultades"].append(f)
    for p in PROGRAMAS_OFICIALES:
        libro["Programas"].append(p)

    # 3. Oferta de materias (copia en memoria para ir descontando cupos reales)
    oferta = construir_oferta_materias()
    oferta_por_curso: dict = {}
    for fila in oferta:
        oferta_por_curso.setdefault((fila[5], fila[1]), []).append(fila)

    # 4. Estudiantes (los 3 integrantes primero)
    estudiantes = [
        ["1001", "Alejandro", "Campo Solarte", "alejandro.campo@campusucc.edu.co", "3158901234", 6, "P01"],
        ["1002", "Juan Esteban", "Cardona Calambas", "juan.cardona@campusucc.edu.co", "3176543210", 6, "P01"],
        ["1003", "Diego Fernando", "Gómez Rivera", "diego.gomez@campusucc.edu.co", "3201122334", 6, "P01"],
    ]

    programas_codigos = [p[0] for p in PROGRAMAS_OFICIALES]
    semestres_max = {p[0]: p[3] for p in PROGRAMAS_OFICIALES}

    for i in range(4, total_estudiantes + 1):
        cedula = f"{1000 + i}"
        nom = NOMBRES_POOL[(i - 4) % len(NOMBRES_POOL)]
        ape = APELLIDOS_POOL[(i - 4) % len(APELLIDOS_POOL)]
        prog = programas_codigos[(i - 1) % len(programas_codigos)]
        sem = min(((i * 3) % 10) + 1, semestres_max[prog])
        nom_limpio = _quitar_tildes(nom.lower().split()[0])
        ape_limpio = _quitar_tildes(ape.lower().split()[0])
        correo = f"{nom_limpio}.{ape_limpio}{i}@campusucc.edu.co"
        tel = f"31{azar.randint(0, 9)}{azar.randint(1000000, 9999999)}"
        estudiantes.append([cedula, nom, ape, correo, tel, sem, prog])

    for e in estudiantes:
        libro["Estudiantes"].append(e)

    # 5. Matrículas y notas
    matriculas = []
    notas = []
    periodo = config.PERIODO_ACTUAL_DEFECTO
    notas_base_integrantes = {
        "1001": [4.8, 4.6, 5.0],
        "1002": [4.5, 4.7, 4.9],
        "1003": [4.6, 4.8, 4.7],
    }

    for e in estudiantes:
        ced, prog, sem = e[0], e[6], int(e[5])
        grupos = _elegir_grupos_sin_cruce(prog, sem, oferta_por_curso, azar)

        for idx_m, grupo in enumerate(grupos):
            grupo[3] -= 1  # Descuento real de cupo en el grupo elegido
            id_matricula = f"MAT{ced}_{grupo[0]}"
            fecha = "2026-09-10" if ced in notas_base_integrantes else "2026-09-15"
            matriculas.append([id_matricula, ced, grupo[0], grupo[2], periodo, fecha, "Activa"])

            if ced in notas_base_integrantes:
                n_base = notas_base_integrantes[ced]
                var = (idx_m * 0.1) - 0.1
                n1 = min(5.0, max(3.5, round(n_base[0] + var, 2)))
                n2 = min(5.0, max(3.5, round(n_base[1] - var, 2)))
                n3 = min(5.0, max(3.5, round(n_base[2] + (var / 2), 2)))
            elif azar.random() < 0.85:
                n1 = round(azar.uniform(2.8, 4.9), 2)
                n2 = round(azar.uniform(2.7, 4.9), 2)
                n3 = round(azar.uniform(3.0, 5.0), 2)
            else:
                continue

            definitiva = round((n1 + n2 + n3) / 3.0, 2)
            estado = "Aprobado" if definitiva >= config.NOTA_APROBATORIA else "Reprobado"
            notas.append([f"NOT_{id_matricula}", id_matricula, grupo[0], n1, n2, n3, definitiva, estado])

    for fila_m in oferta:
        libro["Materias"].append(fila_m)
    for mat in matriculas:
        libro["Matriculas"].append(mat)
    for n in notas:
        libro["Notas"].append(n)

    # 6. Usuarios del sistema
    clave_hash = encriptar_password("1234")
    usuarios = [
        ["admin", clave_hash, "Administrador", "Coordinación Académica UCC", "0000"],
        ["profe", clave_hash, "Profesor", "Docente UCC", "9001"],
    ]
    for e in estudiantes:
        usuarios.append([e[0], clave_hash, "Estudiante", f"{e[1]} {e[2]}", e[0]])

    for u in usuarios:
        libro["Usuarios"].append(u)

    libro.save(archivo_salida)
    libro.close()

    resumen = {
        "ruta": str(archivo_salida),
        "facultades": len(FACULTADES_OFICIALES),
        "programas": len(PROGRAMAS_OFICIALES),
        "materias": len(oferta),
        "asignaturas_unicas": len({codigo_base_materia(f[0]) for f in oferta}),
        "estudiantes": len(estudiantes),
        "matriculas": len(matriculas),
        "notas": len(notas),
        "usuarios": len(usuarios),
    }

    if not silencioso:
        print("==================================================================")
        print(f" [OK] ARCHIVO EXCEL MASIVO GENERADO: {archivo_salida}")
        print("==================================================================")
        print(f" * Facultades:   {resumen['facultades']} | Programas: {resumen['programas']}")
        print(f" * Asignaturas:  {resumen['asignaturas_unicas']} ({resumen['materias']} grupos con horario)")
        print(f" * Estudiantes:  {resumen['estudiantes']} | Matriculas: {resumen['matriculas']} | Notas: {resumen['notas']}")
        print("==================================================================")

    return resumen


def poblar_base_de_datos(silencioso: bool = False):
    """
    Asegura que el archivo de base de datos exista al iniciar la aplicación,
    pero SIN poblarlo automáticamente para que la primera visita sea en blanco.
    """
    libro = db.crear_archivo_excel()
    libro.close()


if __name__ == "__main__":
    ruta_def = config.CARPETA_PLANTILLAS_EXCEL / "Universidad_Completa_50_Estudiantes_UCC.xlsx"
    generar_simulacion_masiva_universidad(total_estudiantes=50, silencioso=False, ruta_destino=ruta_def)
